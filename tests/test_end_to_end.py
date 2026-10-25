"""Comprehensive automated verification suite for LinkBERT Pro."""

import os
import shutil
import unittest
from pathlib import Path

from lb.apply import LinkApplier
from lb.discovery import HybridDiscoveryEngine
from lb.external import ExternalLinkEngine
from lb.graph import LinkGraph
from lb.site import build


class TestLinkBertPro(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.root = Path(__file__).resolve().parent
        cls.fixture_dir = cls.root / "fixture"
        cls.test_index = cls.root / ".test_index"
        try:
            import sentence_transformers
            has_embed = True
        except ImportError:
            has_embed = False
        cls.idx = build(str(cls.fixture_dir), cls.test_index, embed=has_embed)

    @classmethod
    def tearDownClass(cls):
        if cls.test_index.exists():
            shutil.rmtree(cls.test_index)

    def test_01_indexing_and_pages(self):
        self.assertEqual(len(self.idx.pages), 6)
        keys = {p.key for p in self.idx.pages}
        self.assertIn("/", keys)
        self.assertIn("/senior-pets/", keys)
        self.assertIn("/accident-only-cover/", keys)
        self.assertIn("/routine-care/", keys)
        self.assertIn("/compare-pet-insurance/", keys)
        self.assertIn("/blog/choosing-a-vet/", keys)

    def test_02_graph_and_pagerank(self):
        graph = LinkGraph(self.idx)
        metrics = graph.analyze()
        self.assertIn("/compare-pet-insurance/", metrics.page_rank)
        # Compare page should have higher PR because homepage links to it
        self.assertGreater(metrics.page_rank["/compare-pet-insurance/"], 1.0)
        # 4 orphan pages detected
        self.assertEqual(len(metrics.orphan_pages), 4)

    def test_03_hybrid_discovery(self):
        engine = HybridDiscoveryEngine(self.idx, predictor=None)
        senior_page = self.idx.by_key["/senior-pets/"]
        opps = engine.find_direct_mentions(senior_page, set())

        # Must discover direct mentions of accident-only and routine-care
        targets = {o.target_key for o in opps}
        self.assertIn("/accident-only-cover/", targets)
        self.assertIn("/routine-care/", targets)

    def test_04_html_apply_and_rollback(self):
        temp_dir = self.root / "temp_apply_test"
        shutil.copytree(self.fixture_dir, temp_dir)

        try:
            temp_idx = build(str(temp_dir), temp_dir / ".idx", embed=False)
            engine = HybridDiscoveryEngine(temp_idx, predictor=None)
            senior_page = temp_idx.by_key["/senior-pets/"]
            opps = engine.find_direct_mentions(senior_page, set())

            applier = LinkApplier(dry_run=False, create_backup=True)
            res = applier.apply({senior_page.source: opps})
            self.assertEqual(len(res), 1)
            self.assertGreater(res[0].applied_count, 0)

            # Check that file now contains link
            modified_text = Path(senior_page.source).read_text(encoding="utf-8")
            self.assertIn('<a href="/accident-only-cover/">', modified_text)

            # Rollback
            restored = LinkApplier.rollback(temp_dir)
            self.assertGreater(len(restored), 0)
            restored_text = Path(senior_page.source).read_text(encoding="utf-8")
            self.assertNotIn('<a href="/accident-only-cover/">', restored_text)
        finally:
            if temp_dir.exists():
                shutil.rmtree(temp_dir)

    def test_05_external_engine(self):
        ext = ExternalLinkEngine()
        home = self.idx.by_key["/"]
        # Home has no external claims by default
        opps = ext.analyze_page(home)
        self.assertIsInstance(opps, list)

    def test_06_topic_clusters(self):
        from lb.intelligence import detect_topic_clusters
        clusters_hier = detect_topic_clusters(self.idx, method="url_hierarchy")
        self.assertGreaterEqual(len(clusters_hier), 4)

        clusters_louvain = detect_topic_clusters(self.idx, method="louvain")
        self.assertGreaterEqual(len(clusters_louvain), 4)
        # Verify hub page assignment
        for c in clusters_louvain:
            self.assertTrue(c.hub_page in self.idx.by_key)

    def test_07_anchor_diversity(self):
        from lb.intelligence import analyze_anchor_diversity
        # Simulate proposed links
        proposed = [
            ("/senior-pets/", "/accident-only-cover/", "accident-only cover"),
            ("/", "/accident-only-cover/", "accident-only cover"),
            ("/blog/choosing-a-vet/", "/accident-only-cover/", "accident-only cover"),
        ]
        profiles = analyze_anchor_diversity(self.idx, proposed)
        self.assertGreater(len(profiles), 0)
        profile_acc = next(p for p in profiles if p.target_key == "/accident-only-cover/")
        # All 3 anchors are identical -> entropy should be 0 and flagged as over-optimized
        self.assertEqual(profile_acc.entropy, 0.0)
        self.assertTrue(profile_acc.is_over_optimized)

    def test_08_markdown_apply_and_rollback(self):
        temp_md = self.root / "temp_test.md"
        temp_md.write_text("# Guide\n\nFor more details check accident-only cover here.\n", encoding="utf-8")
        try:
            senior_page = self.idx.by_key["/senior-pets/"]
            engine = HybridDiscoveryEngine(self.idx, predictor=None)
            opps = engine.find_direct_mentions(senior_page, set())
            app_opps = [o for o in opps if o.target_key == "/accident-only-cover/"]

            applier = LinkApplier(dry_run=False, create_backup=True)
            res = applier.apply_to_markdown(temp_md, app_opps)
            self.assertEqual(res.applied_count, 1)
            modified = temp_md.read_text(encoding="utf-8")
            self.assertIn("[accident-only cover](/accident-only-cover/)", modified)

            # Rollback
            restored = LinkApplier.rollback(self.root)
            self.assertIn(str(temp_md), restored)
            rolled_back = temp_md.read_text(encoding="utf-8")
            self.assertNotIn("](/accident-only-cover/)", rolled_back)
        finally:
            if temp_md.exists():
                temp_md.unlink()
            bak = temp_md.with_suffix(temp_md.suffix + ".bak")
            if bak.exists():
                bak.unlink()

    def test_09_cannibalization_detection(self):
        from lb.browser import parse_search_console_csv
        from lb.intelligence import detect_cannibalization
        csv_text = (
            "Top queries,Clicks,Impressions,CTR,Position,Top pages\n"
            "pet insurance,100,2000,5.0%,3.0,https://example.com/page1\n"
            "pet insurance,50,1500,3.3%,7.5,https://example.com/page2\n"
        )
        rows = parse_search_console_csv(csv_text)
        self.assertEqual(len(rows), 2)
        issues = detect_cannibalization(rows, min_impressions=10)
        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0].query, "pet insurance")
        self.assertIn("https://example.com/page1", issues[0].recommendation)

    def test_10_serp_conflict_guard(self):
        from lb.browser import SERPConflict, SERPResult, SearchConsoleRow, check_serp_conflicts
        ext_domains = {"competitor.com", "gov.au"}
        sc_rows = [SearchConsoleRow(query="pet insurance", page="/", clicks=10, impressions=100, ctr=0.1, position=3.0)]
        serp = {
            "pet insurance": [
                SERPResult(title="Competitor", url="https://competitor.com/best-pet-insurance", domain="competitor.com", position=4, snippet="..."),
                SERPResult(title="Wiki", url="https://en.wikipedia.org/wiki/Pet_insurance", domain="wikipedia.org", position=1, snippet="..."),
            ]
        }
        conflicts = check_serp_conflicts(ext_domains, sc_rows, serp)
        self.assertEqual(len(conflicts), 1)
        self.assertEqual(conflicts[0].external_domain, "competitor.com")
        self.assertEqual(conflicts[0].risk_level, "high")

    def test_11_external_link_engine(self):
        from lb.external import ExternalLinkEngine
        ext_engine = ExternalLinkEngine()
        all_ext = []
        for p in self.idx.pages:
            all_ext.extend(ext_engine.analyze_page(p))
        self.assertIsInstance(all_ext, list)
        for o in all_ext:
            self.assertIn(o.trust_tier, (1, 2, 3, 4))
            self.assertTrue(len(o.anchor) > 1)
            self.assertIn("noopener", o.recommended_rel)

    def test_12_multi_tab_excel_export(self):
        import openpyxl
        from lb.report import write_all
        from lb.discovery import HybridDiscoveryEngine
        eng = HybridDiscoveryEngine(self.idx)
        senior_page = self.idx.by_key["/senior-pets/"]
        opps = eng.find_direct_mentions(senior_page, set())
        stats = write_all(self.test_index, self.idx, opps)
        self.assertIn("excel", stats)
        wb = openpyxl.load_workbook(stats["excel"], read_only=True)
        self.assertIn("Internal_Linking_Matrix", wb.sheetnames)
        self.assertIn("External_Authority_Citations", wb.sheetnames)
        self.assertIn("GSC_Striking_Distance", wb.sheetnames)
        self.assertIn("Graph_Topology_Audit", wb.sheetnames)
        wb.close()

    def test_13_specialist_audit_gate(self):
        from tests.verify_deliverable import audit_workbook
        stats = audit_workbook(self.test_index / "link_architecture_master.xlsx")
        self.assertEqual(stats["status"], "PASS")
        self.assertEqual(stats["total_defects"], 0)


if __name__ == "__main__":
    unittest.main()
